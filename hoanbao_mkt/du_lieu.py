"""Đọc file Excel dữ liệu nền và làm sạch: bỏ mọi thông tin chưa chắc chắn."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

CAN_BO_SUNG = "[cần bổ sung]"
# Đoạn chứa các cụm này bị loại, vì file ghi rõ là chưa xác nhận.
DAU_HIEU_CHUA_CHAC = (CAN_BO_SUNG, "xác nhận")
# Ghi chú nguồn kiểu "(theo nhãn)", "(theo website)" không đưa vào bài đăng.
_GHI_CHU_NGUON = re.compile(r"\s*\(theo [^)]*\)", re.IGNORECASE)
_RONG = {"", "—", "-", "–"}


def lam_sach(gia_tri) -> str | None:
    """Trả về chuỗi đã làm sạch, hoặc None nếu ô không có thông tin dùng được.

    Ô có nhiều ý ngăn bằng ";" thì chỉ bỏ ý chưa chắc chắn, giữ lại phần còn lại.
    """
    if gia_tri is None:
        return None
    text = str(gia_tri).strip()
    giu_lai = []
    for doan in text.split(";"):
        doan = doan.strip()
        if any(d in doan.lower() for d in DAU_HIEU_CHUA_CHAC):
            continue
        doan = _GHI_CHU_NGUON.sub("", doan).strip()
        if doan.rstrip(":").strip() in _RONG or doan.endswith(":"):
            continue
        giu_lai.append(doan)
    ket_qua = "; ".join(giu_lai)
    return ket_qua or None


@dataclass
class SanPham:
    ma: str
    ten: str
    nhom: str
    thu_muc: str | None
    chu_luc: bool
    thong_tin: dict[str, str] = field(default_factory=dict)

    @property
    def la_phu_gia(self) -> bool:
        return self.nhom.lower().startswith("phụ gia")


@dataclass
class NhomKhach:
    ten: str
    mua_qua: str
    ma_san_pham: list[str]
    thong_tin: dict[str, str] = field(default_factory=dict)


@dataclass
class KhungLich:
    thu: int  # 0 = Thứ 2 ... 6 = Chủ nhật
    ten_thu: str
    loai_bai: str
    muc_tieu: str
    goi_y: str
    nguon_media: str
    gio_dang: str  # "HH:MM"


@dataclass
class DuLieuNen:
    cong_ty: list[str]
    san_pham: list[SanPham]
    nhom_khach: list[NhomKhach]
    lich: list[KhungLich]

    def tim_san_pham(self, ma: str) -> SanPham | None:
        return next((sp for sp in self.san_pham if sp.ma.upper() == ma.strip().upper()), None)


_THU = {"thứ 2": 0, "thứ 3": 1, "thứ 4": 2, "thứ 5": 3, "thứ 6": 4, "thứ 7": 5, "chủ nhật": 6}
_GIO = re.compile(r"^\d{1,2}:\d{2}$")


def _doc_bang(ws, cot_bo_qua: set[str]) -> list[dict[str, str | None]]:
    """Dòng 1 là ghi chú, dòng 2 là tiêu đề, dữ liệu từ dòng 3."""
    rows = list(ws.iter_rows(values_only=True))
    tieu_de = [str(c).strip() if c else "" for c in rows[1]]
    bang = []
    for row in rows[2:]:
        if not any(c not in (None, "") for c in row):
            continue
        bang.append({t: row[i] for i, t in enumerate(tieu_de) if t and t not in cot_bo_qua})
    return bang


def _thong_tin_sach(dong: dict, bo: set[str]) -> dict[str, str]:
    ket_qua = {}
    for k, v in dong.items():
        if k in bo:
            continue
        sach = lam_sach(v)
        if sach:
            ket_qua[k] = sach
    return ket_qua


def doc_du_lieu(duong_dan: str | Path, cot_bo_qua: list[str] | None = None) -> DuLieuNen:
    wb = openpyxl.load_workbook(duong_dan, data_only=True, read_only=True)
    bo_qua = set(cot_bo_qua or [])

    # Thông tin công ty: các dòng sau "THÔNG TIN CÔNG TY", bỏ dòng "LƯU Ý".
    cong_ty: list[str] = []
    bat_dau = False
    for (o,) in wb["Hướng dẫn"].iter_rows(values_only=True, max_col=1):
        text = str(o or "").strip()
        if text.upper().startswith("THÔNG TIN CÔNG TY"):
            bat_dau = True
            continue
        if bat_dau and text and not text.upper().startswith("LƯU Ý"):
            sach = lam_sach(text)
            if sach:
                cong_ty.append(sach)

    san_pham = []
    for d in _doc_bang(wb["Sản phẩm"], set()):
        ma = str(d.get("Mã SP") or "").strip()
        if not ma:
            continue
        san_pham.append(SanPham(
            ma=ma,
            ten=str(d.get("Tên sản phẩm") or ma).strip(),
            nhom=str(d.get("Nhóm keo") or "").strip(),
            thu_muc=lam_sach(d.get("Thư mục ảnh/video trong kho")),
            chu_luc=str(d.get("Chủ lực?") or "").strip().lower() == "có",
            thong_tin=_thong_tin_sach(d, bo_qua | {"Mã SP", "Thư mục ảnh/video trong kho", "Chủ lực?"}),
        ))

    nhom_khach = []
    for d in _doc_bang(wb["Khách hàng mục tiêu"], set()):
        ten = str(d.get("Nhóm khách") or "").strip()
        if not ten:
            continue
        sp_text = str(d.get("Sản phẩm phù hợp") or "")
        ma_sp = [m.strip() for m in sp_text.split(",") if m.strip() and m.strip().lower() != "toàn bộ"]
        nhom_khach.append(NhomKhach(
            ten=ten,
            mua_qua=str(d.get("Mua qua") or "").strip(),
            ma_san_pham=ma_sp,
            thong_tin=_thong_tin_sach(d, bo_qua),
        ))

    lich = []
    for d in _doc_bang(wb["Lịch chủ đề"], set()):
        ten_thu = str(d.get("Ngày") or "").strip()
        gio = str(d.get("Giờ đăng") or "").strip()
        if ten_thu.lower() not in _THU or not _GIO.match(gio):
            continue  # ví dụ Chủ nhật "Không đăng"
        lich.append(KhungLich(
            thu=_THU[ten_thu.lower()],
            ten_thu=ten_thu,
            loai_bai=str(d.get("Loại bài") or "").strip(),
            muc_tieu=str(d.get("Mục tiêu") or "").strip(),
            goi_y=str(d.get("Gợi ý nội dung") or "").strip(),
            nguon_media=str(d.get("Ảnh/video lấy từ") or "").strip(),
            gio_dang=gio.zfill(5),
        ))
    wb.close()
    return DuLieuNen(cong_ty=cong_ty, san_pham=san_pham, nhom_khach=nhom_khach, lich=lich)
